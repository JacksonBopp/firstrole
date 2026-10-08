"""The tracker: one CSV row per job, the single source of truth.

Every row has a stable short `id` derived from the job URL, so updates target exactly
one job (searching by text can match several). Notes are append-only: each update adds
a dated line instead of overwriting history.
"""
from __future__ import annotations

import csv
import hashlib
from collections import Counter
from datetime import date
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from . import paths

COLUMNS = ["id", "date_found", "company", "title", "category", "location", "url",
           "resume", "status", "date_applied", "follow_up", "notes"]

# Canonical statuses, in pipeline order. Free-text suffixes ("Skipped - no visa") are allowed.
STATUSES = ["Found", "Queued", "Applied", "Assessment", "Interviewing", "Offer",
            "Rejected", "Withdrawn", "Skipped", "Blocked"]
ACTIVE = {"Applied", "Assessment", "Interviewing", "Offer"}


def normalize_url(url: str) -> str:
    """Drop tracking query strings and trailing slashes so the same posting maps to one row."""
    parts = urlsplit(url.strip())
    keep = "&".join(q for q in parts.query.split("&") if q and q.split("=")[0] in {"jobId", "gh_jid", "reqId", "id"})
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path.rstrip("/"), keep, ""))


def job_id(url: str) -> str:
    return hashlib.sha1(normalize_url(url).encode()).hexdigest()[:8]


def load(path: Path | None = None) -> list[dict]:
    path = path or paths.tracker_csv()
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def save(rows: list[dict], path: Path | None = None) -> None:
    path = path or paths.tracker_csv()
    extra = [c for r in rows for c in r if c not in COLUMNS]
    cols = COLUMNS + list(dict.fromkeys(extra))           # keep unknown columns users add
    tmp = path.with_suffix(".tmp")
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, restval="")
        w.writeheader()
        w.writerows(rows)
    tmp.replace(path)                                       # atomic on Windows and POSIX


def get(rows: list[dict], key: str) -> dict:
    """Look up one row by id or URL. Raises if missing."""
    k = key.strip()
    for r in rows:
        if r["id"] == k or r["url"] == k or normalize_url(r["url"]) == normalize_url(k):
            return r
    raise KeyError(f"No tracked job with id/url {key!r}")


def find(rows: list[dict], text: str) -> list[dict]:
    t = text.lower()
    return [r for r in rows if t in f"{r['company']} {r['title']} {r['url']} {r['id']}".lower()]


def add(rows: list[dict], *, company: str, title: str, url: str, status: str = "Found",
        category: str = "", location: str = "", resume: str = "", applied: str = "", note: str = "") -> dict:
    jid = job_id(url)
    if any(r["id"] == jid for r in rows):
        raise ValueError(f"Already tracked as {jid}; use update instead")
    row = dict.fromkeys(COLUMNS, "")
    row.update(id=jid, date_found=date.today().isoformat(), company=company.strip(), title=title.strip(),
               category=category, location=location, url=url.strip(), resume=resume, status=status,
               date_applied=applied, notes=_note("", note))
    rows.append(row)
    return row


def update(rows: list[dict], key: str, *, status: str | None = None, applied: str | None = None,
           follow_up: str | None = None, resume: str | None = None, note: str | None = None) -> dict:
    r = get(rows, key)
    for field, val in (("status", status), ("date_applied", applied), ("follow_up", follow_up), ("resume", resume)):
        if val is not None:
            r[field] = val
    if note:
        r["notes"] = _note(r["notes"], note)
    return r


def _note(existing: str, new: str) -> str:
    if not new:
        return existing
    line = f"{date.today().isoformat()}: {new.strip()}"
    return f"{existing} | {line}" if existing else line


def company_counts(rows: list[dict]) -> Counter:
    """Applications per company (only rows that were actually submitted)."""
    return Counter(r["company"].strip().lower() for r in rows if r["date_applied"])


def summary(rows: list[dict]) -> list[tuple[str, int]]:
    return Counter(r["status"].split(" - ")[0] for r in rows).most_common()


def export_xlsx(rows: list[dict], out: Path) -> int:
    """Write a filtered, styled spreadsheet of everything acted on. Needs `pip install job-search-os[xlsx]`."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as e:
        raise SystemExit("Spreadsheet export needs openpyxl: pip install 'job-search-os[xlsx]'") from e
    shown = sorted((r for r in rows if r["status"] != "Found"), key=lambda r: r["date_applied"] or "", reverse=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Applications"
    cols = ["company", "title", "location", "status", "date_applied", "follow_up", "resume", "url", "notes"]
    ws.append([c.replace("_", " ").title() for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1B1F24")
    for r in shown:
        ws.append([r.get(c, "") for c in cols])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    return len(shown)
